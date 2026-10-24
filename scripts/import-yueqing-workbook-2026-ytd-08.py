#!/usr/bin/env python3
"""Import the supplied 2026 Jan-Aug ranking workbook into the website dataset."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


BAND = re.compile(r"^(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)\s*万(?:美元|美金)?$")


def normalize_name(name: str) -> str:
    return re.sub(r"[\s\u3000（）()·•,，.。\-—_]+", "", name)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--aliases", type=Path, required=True)
    parser.add_argument("--period-id", default="2026-ytd-08")
    parser.add_argument("--period-label", default="2026年1—8月")
    parser.add_argument("--as-of", default="2026-08-31")
    return parser.parse_args()


def read_workbook(path: Path) -> list[dict]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    header = next(sheet.iter_rows(min_row=1, max_col=3, values_only=True))
    if tuple(value.strip() if isinstance(value, str) else value for value in header) != (
        "序号",
        "企业名称",
        "累计出口美元区间",
    ):
        raise ValueError(f"Unexpected workbook header: {header!r}")

    records = []
    for excel_row, (raw_rank, raw_name, raw_band) in enumerate(
        sheet.iter_rows(min_row=2, max_col=3, values_only=True),
        start=2,
    ):
        if raw_rank is None and raw_name is None and raw_band is None:
            continue
        if isinstance(raw_rank, bool) or not isinstance(raw_rank, (int, float)):
            raise ValueError(f"Excel row {excel_row}: invalid rank {raw_rank!r}")
        rank = int(raw_rank)
        if rank != raw_rank:
            raise ValueError(f"Excel row {excel_row}: non-integer rank {raw_rank!r}")
        name = str(raw_name or "").strip()
        if not name:
            raise ValueError(f"Excel row {excel_row}: blank company name")
        source_band = str(raw_band or "").strip()
        band = source_band or "100万美元以下"
        records.append(
            {
                "rank": rank,
                "company_name": name,
                "band_label": band,
                "source_band_blank": not source_band,
                "source_file": path.name,
                "source_line": excel_row,
            }
        )

    ranks = [record["rank"] for record in records]
    if ranks != sorted(ranks):
        raise ValueError("Ranks must remain sorted")
    if len(ranks) != len(set(ranks)):
        raise ValueError("Ranks must be unique")
    expected = list(range(1, ranks[-1] + 1))
    missing_ranks = sorted(set(expected) - set(ranks))
    if missing_ranks != [43]:
        raise ValueError(f"Unexpected source rank gaps: {missing_ranks}")
    return records


def build_bands(records: list[dict]) -> list[dict]:
    labels = list(dict.fromkeys(record["band_label"] for record in records))
    bands = []
    for index, label in enumerate(labels, start=1):
        match = BAND.fullmatch(label)
        bands.append(
            {
                "id": f"b{index:02d}",
                "label": label,
                "min_usd_10k": float(match.group(1)) if match else None,
                "max_usd_10k": float(match.group(2)) if match else None,
                "count": sum(record["band_label"] == label for record in records),
            }
        )
    return bands


def write_json(path: Path, value: object) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )


def main() -> None:
    args = parse_args()
    records = read_workbook(args.workbook)
    companies_path = args.data_dir / "companies.json"
    companies = json.loads(companies_path.read_text(encoding="utf-8"))
    original_company_count = len(companies)
    explicit_aliases = json.loads(args.aliases.read_text(encoding="utf-8"))

    by_name: dict[str, str] = {}
    normalized: dict[str, set[str]] = defaultdict(set)
    by_id = {company["company_id"]: company for company in companies}
    for company in companies:
        for name in [company["company_name"], *company.get("aliases", [])]:
            by_name[name] = company["company_id"]
            normalized[normalize_name(name)].add(company["company_id"])

    for alias, canonical in explicit_aliases.items():
        if canonical not in by_name:
            raise ValueError(f"Alias target is missing from registry: {canonical}")
        company_id = by_name[canonical]
        by_name[alias] = company_id
        normalized[normalize_name(alias)].add(company_id)
        if alias not in by_id[company_id]["aliases"]:
            by_id[company_id]["aliases"].append(alias)

    name_counts = Counter(record["company_name"] for record in records)
    next_id = max(int(company["company_id"][2:]) for company in companies) + 1
    matched_records = []
    for record in records:
        source_name = record["company_name"]
        company_id = by_name.get(source_name)
        if company_id is None:
            candidates = normalized[normalize_name(source_name)]
            if len(candidates) == 1:
                company_id = next(iter(candidates))
                if source_name not in by_id[company_id]["aliases"]:
                    by_id[company_id]["aliases"].append(source_name)
            elif candidates:
                raise ValueError(f"Ambiguous normalized company name: {source_name}")
            else:
                company_id = f"YQ{next_id:06d}"
                next_id += 1
                company = {
                    "company_id": company_id,
                    "company_name": source_name,
                    "aliases": [],
                    "needs_review": name_counts[source_name] > 1,
                }
                companies.append(company)
                by_id[company_id] = company
                by_name[source_name] = company_id
                normalized[normalize_name(source_name)].add(company_id)
        if name_counts[source_name] > 1:
            by_id[company_id]["needs_review"] = True
        matched_records.append({**record, "company_id": company_id})

    bands = build_bands(records)
    band_ids = {band["label"]: band["id"] for band in bands}
    output_records = [
        {
            "rank": record["rank"],
            "company_id": record["company_id"],
            "company_name": record["company_name"],
            "band_id": band_ids[record["band_label"]],
            "band_label": record["band_label"],
            "source_file": record["source_file"],
            "source_line": record["source_line"],
            "source_duplicate": name_counts[record["company_name"]] > 1,
        }
        for record in matched_records
    ]
    period = {
        "id": args.period_id,
        "label": args.period_label,
        "type": "year_to_date",
        "as_of": args.as_of,
        "metric": "exports",
        "precision": "range",
    }
    data = {
        "period": period,
        "source_note": "排名来自提供的累计出口区间数据；仅展示累计排名与出口额区间，不推算精确出口额。",
        "validation": {
            "source_rank_gaps": [43],
            "source_blank_band_records": sum(record["source_band_blank"] for record in records),
        },
        "bands": bands,
        "records": output_records,
    }

    manifest_path = args.data_dir / "periods.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    period_meta = {
        **period,
        "file": f"{args.period_id}.json",
        "record_count": len(output_records),
    }
    manifest["periods"] = [
        item for item in manifest["periods"] if item["id"] != args.period_id
    ] + [period_meta]
    manifest["periods"].sort(key=lambda item: item["as_of"], reverse=True)
    manifest["latest"] = manifest["periods"][0]["id"]

    write_json(args.data_dir / f"{args.period_id}.json", data)
    write_json(companies_path, companies)
    write_json(manifest_path, manifest)
    print(
        json.dumps(
            {
                "records": len(output_records),
                "companies": len(companies),
                "new_companies": len(companies) - original_company_count,
                "bands": len(bands),
                "missing_ranks": data["validation"]["source_rank_gaps"],
                "source_blank_band_records": data["validation"]["source_blank_band_records"],
                "duplicate_records": sum(record["source_duplicate"] for record in output_records),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
